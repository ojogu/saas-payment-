import random
import string
from io import BytesIO

import jwt
import requests
from flask import Blueprint, current_app, jsonify, request
from openpyxl import load_workbook

from src.model import (
    Accounts,
    Clearance,
    Notification_Setting,
    PassivePayment,
    Payment,
    User,
    db,
)

ALLOWED_EXTENSIONS = {"xlsx"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def generate_payment_reference(prefix="ref", length=10):
    random_part = "".join(
        random.choices(string.ascii_lowercase + string.digits, k=length)
    )
    return f"{prefix}-{random_part}"


payment_routes = Blueprint("payment", __name__)


@payment_routes.route("/passive_payment/initiate", methods=["POST"])
def passive_payment():
    try:
        SECRET_KEY = "balablu-01101"
        decoded = jwt.decode(
            request.headers.get("Authorization"), SECRET_KEY, algorithms=["HS256"]
        )
        user = decoded
    except:
        return jsonify({"message": "Invalid auth token"})
    if user.get("role") not in ["STUDENT", "DATA"]:
        return jsonify({"message": "Unauthorized Access"}), 401
    data = request.get_json()
    student_id = user.get("id")
    payment_ref = data.get("ref")
    platform = data.get("platform")
    item = data.get("item")
    complete = data.get("complete")

    amount = 0
    if user.get("role") == "DATA":
        student_id = data.get("student_id")
    if not payment_ref:
        return jsonify({"message": "No Payment Reference"})
    if not platform:
        return jsonify({"message": "No Plaform selected"})

    if complete:
        clearance = Clearance.query.get(item)
        item = clearance.clearance_item_id

    payment = Payment.query.filter_by(payment_ref=payment_ref).first()
    if payment:
        return jsonify({"message": "This receipt has already been used"})

    passive_payment = PassivePayment.query.filter_by(payment_ref=payment_ref).first()
    if passive_payment and passive_payment.clearance_item_id == item:
        amount = passive_payment.amount
        new_payment = Payment(
            user_id=student_id, payment_ref=payment_ref, amount=amount
        )
        db.session.add(new_payment)
        db.session.commit()
    else:
        return jsonify({"message": "Incorrect Reference ID"})
    link = f"{payment_ref}"
    return jsonify(
        {
            "status": "success",
            "message": "Payment Verified Successfully",
            "link": link,
            "payment_id": new_payment.id,
        }
    )


@payment_routes.route("/system_payment/verify/<id>", methods=["GET"])
def system_payment_verify(id):
    payment = Payment.query.get(id)
    headers = {
        "Authorization": f"Bearer {current_app.config['PAYSTACK_SECRET_KEY']}",
        "Content-Type": "application/json",
    }

    response = requests.get(
        f"https://api.paystack.co/transaction/verify/{payment.payment_ref}",
        headers=headers,
    )
    result = response.json()
    if result["status"] and result["data"]["status"] == "success":
        payment.status = "paid"
        db.session.commit()
        return jsonify(
            {
                "status": "success",
                "message": "System Payment made successfully you can begin using the system",
            }
        )
    else:
        return jsonify({"message": "No payment made"})


@payment_routes.route("/notification_payment/initiate", methods=["POST"])
def notification_payment():
    try:
        SECRET_KEY = "balablu-01101"
        decoded = jwt.decode(
            request.headers.get("Authorization"), SECRET_KEY, algorithms=["HS256"]
        )
        user = decoded
    except:
        return jsonify({"message": "Invalid auth token"})
    if user.get("role") not in ["STUDENT"]:
        return jsonify({"message": "Unauthorized Access"}), 401
    student_id = user.get("id")
    payment_ref = generate_payment_reference()
    notify_setting = Notification_Setting.query.filter_by(
        organization_id=user.get("organization_id")
    ).first()
    new_payment = Payment(
        user_id=student_id,
        payment_ref=payment_ref,
        system_payment=True,
        amount=notify_setting.price,
    )
    account = Accounts.query.filter_by(default_account=True).first()
    db.session.add(new_payment)
    db.session.commit()
    link = f"{payment_ref}"
    return jsonify(
        {
            "message": "Payment Initiated Successfully",
            "link": link,
            "payment_id": new_payment.id,
            "subaccount": account.subaccount,
            "amount": notify_setting.price,
        }
    )


@payment_routes.route("/notification_payment/verify/<id>", methods=["GET"])
def notification_payment_verify(id):
    payment = Payment.query.get(id)
    headers = {
        "Authorization": f"Bearer {current_app.config['PAYSTACK_SECRET_KEY']}",
        "Content-Type": "application/json",
    }

    response = requests.get(
        f"https://api.paystack.co/transaction/verify/{payment.payment_ref}",
        headers=headers,
    )
    result = response.json()
    if result["status"] and result["data"]["status"] == "success":
        payment.status = "paid"
        db.session.commit()
        return jsonify(
            {
                "status": "success",
                "message": "Notification Payment made successfully you can begin using the system",
            }
        )
    else:
        return jsonify({"message": "No payment made"})


@payment_routes.route("/upload/passive_payments", methods=["POST"])
def upload_payments():
    try:
        SECRET_KEY = "balablu-01101"
        print(request.headers.get("Authorization"))
        decoded = jwt.decode(
            request.headers.get("Authorization"), SECRET_KEY, algorithms=["HS256"]
        )
        user = decoded
    except Exception:
        return jsonify({"message": "Invalid auth token"})
    if user.get("role") not in ["ADMIN", "BURSAR"]:
        return jsonify({"message": "Unauthorized Access"}), 401
    if "file" not in request.files:
        return jsonify({"message": "No file part"}), 400
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"message": "No selected file"}), 400

    if not allowed_file(file.filename):
        return jsonify(
            {"message": "Invalid file type. Only Excel (.xlsx) files are allowed."}
        ), 400

    # Load the Excel file
    try:
        file.stream.seek(0)
        workbook = load_workbook(filename=BytesIO(file.read()))
        sheet = workbook.active  # Get the first sheet
    except Exception as e:
        return jsonify({"message": f"Failed to read Excel file: {str(e)}"}), 400

    users = []
    existing_payments = []
    references = []
    skipped = []
    data = request.form
    item = data.get("item")
    # Iterate through rows, assuming the first row contains headers
    for row in sheet.iter_rows(min_row=2, values_only=True):  # Skip the header row
        payment_ref, matric, amount = row[:3]

        user = User.query.filter_by(matric_number=matric).first()
        existing_payment = PassivePayment.query.filter_by(
            payment_ref=payment_ref
        ).first()

        if payment_ref not in references and not existing_payment:
            new_passive_payment = PassivePayment(
                user_id=user.id,
                payment_ref=payment_ref,
                amount=amount,
                clearance_item_id=item,
            )
            db.session.add(new_passive_payment)
            db.session.commit()

            references.append(payment_ref)

    return jsonify(
        {
            "status": "success",
            "message": f"{len(references)} payments uploaded successfully",
        }
    )
